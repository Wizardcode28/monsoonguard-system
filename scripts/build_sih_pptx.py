import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors (MoES / Agro-Meteorology Palette)
NAVY = RGBColor(15, 23, 42)          # #0F172A - Deep MoES Slate Navy
EMERALD = RGBColor(5, 150, 105)      # #059669 - Agriculture Emerald Green
CYAN_BLUE = RGBColor(2, 132, 199)    # #0284C7 - Atmosphere Cyan Blue
DARK_GRAY = RGBColor(51, 65, 85)     # #334155 - Slate Body Text
CARD_BG = RGBColor(248, 250, 252)    # #F8FAFC
CARD_BORDER = RGBColor(203, 213, 225)# #CBD5E1
AMBER = RGBColor(217, 119, 6)        # #D97706 - Warning Amber

IMG_DIR = r"d:\SIH_2026\monsoonguard-system\docs\images"
LOGO_PATH = os.path.join(IMG_DIR, "sih_bulb_logo_extracted.png")

def set_slide_header(slide, title, subtitle=""):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10.0), Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Georgia"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    if subtitle:
        tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(10.0), Inches(0.5))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.name = "Arial"
        p2.font.size = Pt(15)
        p2.font.bold = True
        p2.font.color.rgb = CYAN_BLUE
        
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(11.9), Inches(0.2), width=Inches(1.05))

# ========================================================
# SLIDE 1: SMART INDIA HACKATHON 2026
# ========================================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_header(s1, "SMART INDIA HACKATHON 2026", "MonsoonGuard AI: Hyper-Local Monsoon Prediction & Agro-Advisory Engine")

# Architecture diagram on left
arch_img = os.path.join(IMG_DIR, "diagram-architecture.png")
if os.path.exists(arch_img):
    s1.shapes.add_picture(arch_img, Inches(0.8), Inches(1.75), width=Inches(6.8))

# Right Table Card
card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.75), Inches(4.7), Inches(5.1))
card1.fill.solid()
card1.fill.fore_color.rgb = CARD_BG
card1.line.color.rgb = CARD_BORDER
card1.line.width = Pt(1.5)

tf1 = card1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_right = Inches(0.25)
tf1.margin_top = Inches(0.25)

s1_data = [
    ("Problem Statement ID", "26086"),
    ("Problem Statement Title", "AI/ML-Based Downscaling of NWP for Block-Level Monsoon Rainfall & Extreme Events"),
    ("Ministry / Organization", "Ministry of Earth Sciences (MoES) / NCMRWF"),
    ("Theme", "Agriculture, FoodTech & Rural Development"),
    ("PS Category", "Software"),
    ("Live Frontend URL", "https://monsoonguard-system.vercel.app/"),
    ("Live Backend API", "https://monsoonguard-system.onrender.com/docs"),
    ("Technical Report (PDF)", "https://drive.google.com/file/d/1fH54z9E90YvwI5SwaRdL1u5-AXliAjqK/view?usp=sharing"),
    ("Demo Video Walkthrough", "https://drive.google.com/file/d/11RcsCXPiOgg9B4bcHfbTABs25oBDbCiB/view?usp=sharing")
]

for idx, (label, val) in enumerate(s1_data):
    p = tf1.paragraphs[0] if idx == 0 else tf1.add_paragraph()
    p.space_after = Pt(8)
    run1 = p.add_run()
    run1.text = f"{label}: "
    run1.font.bold = True
    run1.font.size = Pt(12)
    run1.font.color.rgb = NAVY
    
    run2 = p.add_run()
    run2.text = val
    run2.font.bold = False
    run2.font.size = Pt(11)
    run2.font.color.rgb = DARK_GRAY

# ========================================================
# SLIDE 2: PROBLEM STATEMENT & GROUND CHALLENGES
# ========================================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_header(s2, "Problem Statement & Ground Reality", "Limitations of Coarse Numerical Models for Grassroots Farming")

probs = [
    ("12 km Coarse Grid NWP", "Numerical models (GFS / NCUM) operate on 12 km grid cells. Local micro-climates, rain-shadows, and convective storms are smoothed out over hundreds of square kilometers.", AMBER),
    ("Critical Monsoon Breaks", "Monsoon dry spells (5+ consecutive rainless days during peak vegetative growth) devastate rainfed Kharif crops like Soybean and Paddy if not anticipated 7-10 days early.", NAVY),
    ("Complex Teleconnections", "Equatorial Pacific ENSO, Indian Ocean Dipole, and MJO waves dictate intra-seasonal monsoon bursts and breaks through highly non-linear atmospheric waves.", CYAN_BLUE),
    ("Actionability Gap", "Farmers cannot apply raw millimeter rainfall values. They need direct operational advice: spray windows, sowing delays, and field drainage measures.", EMERALD)
]

for idx, (head, desc, col) in enumerate(probs):
    x = Inches(0.8 + (idx % 2) * 6.0)
    y = Inches(1.8 + (idx // 2) * 2.6)
    c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.3))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = col
    c.line.width = Pt(1.5)
    
    ctf = c.text_frame
    ctf.word_wrap = True
    ctf.margin_left = ctf.margin_right = Inches(0.3)
    ctf.margin_top = Inches(0.25)
    
    p = ctf.paragraphs[0]
    p.text = head
    p.font.name = "Georgia"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = col
    p.space_after = Pt(8)
    
    p2 = ctf.add_paragraph()
    p2.text = desc
    p2.font.name = "Arial"
    p2.font.size = Pt(13)
    p2.font.color.rgb = DARK_GRAY

# ========================================================
# SLIDE 3: PROPOSED INNOVATION & CORE ARCHITECTURE
# ========================================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_header(s3, "Proposed Innovation: Multi-Tier ML Downscaling", "Coupling Planetary Climate Oscillations with 1 km Block-Level Quantile Regressors")

# Screenshot on left
dash_img = os.path.join(IMG_DIR, "02_realtime_dashboard.png")
if os.path.exists(dash_img):
    s3.shapes.add_picture(dash_img, Inches(0.8), Inches(1.75), width=Inches(6.2))

# Right explanation card
c3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.75), Inches(5.2), Inches(5.1))
c3.fill.solid()
c3.fill.fore_color.rgb = CARD_BG
c3.line.color.rgb = CARD_BORDER
c3.line.width = Pt(1.5)

tf3 = c3.text_frame
tf3.word_wrap = True
tf3.margin_left = tf3.margin_right = Inches(0.3)
tf3.margin_top = Inches(0.3)

pts3 = [
    ("Hyper-Local 1 km Downscaling", "LightGBM quantile regressors trained on historical IMD gridded observations and digital elevation models, achieving 50.4% lower RMSE than raw GFS."),
    ("Planetary Teleconnection Assimilation", "Real-time ingestion of ENSO Niño 3.4 SST anomalies, Indian Ocean Dipole (DMI), and 8-phase MJO velocity vectors."),
    ("Probabilistic Quantile Bounds (P10, P50, P90)", "Provides farmers and officials with risk envelopes rather than misleading single point rainfall forecasts."),
    ("Multilingual Agronomic Advisory AI", "Rule-based agro-meteorological engine translating forecasts into crop-stage operations (Soybean, Paddy) in Hindi and English.")
]

for idx, (t, d) in enumerate(pts3):
    p = tf3.paragraphs[0] if idx == 0 else tf3.add_paragraph()
    p.space_after = Pt(10)
    r1 = p.add_run()
    r1.text = f"• {t}\n"
    r1.font.bold = True
    r1.font.size = Pt(14)
    r1.font.color.rgb = NAVY
    r2 = p.add_run()
    r2.text = d
    r2.font.size = Pt(12)
    r2.font.color.rgb = DARK_GRAY

# ========================================================
# SLIDE 4: ML FORECAST & DRY SPELL PREDICTION
# ========================================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_header(s4, "ML Forecast Downscaling & Dry Spell Radar", "16-Day Ensembles with Early Break-Monsoon Warnings")

f_img = os.path.join(IMG_DIR, "03_ml_forecast_downscaling.png")
if os.path.exists(f_img):
    s4.shapes.add_picture(f_img, Inches(0.8), Inches(1.75), width=Inches(7.2))

# Right benchmark card
c4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.75), Inches(4.2), Inches(5.1))
c4.fill.solid()
c4.fill.fore_color.rgb = CARD_BG
c4.line.color.rgb = EMERALD
c4.line.width = Pt(1.5)

tf4 = c4.text_frame
tf4.word_wrap = True
tf4.margin_left = tf4.margin_right = Inches(0.25)
tf4.margin_top = Inches(0.3)

p4 = tf4.paragraphs[0]
p4.text = "Model Evaluation Benchmarks"
p4.font.size = Pt(16)
p4.font.bold = True
p4.font.color.rgb = EMERALD
p4.space_after = Pt(14)

benchmarks = [
    ("RMSE Reduction", "6.18 mm/day (-50.4% vs GFS)"),
    ("MAE Accuracy", "3.72 mm/day (-54.2% error)"),
    ("Dry Spell Break F1", "0.912 (High precision)"),
    ("Heavy Rain ROC-AUC", "0.941 (Accurate alert trigger)"),
    ("Inference Latency", "< 120 ms (Instant API response)")
]

for b_label, b_val in benchmarks:
    p = tf4.add_paragraph()
    p.space_after = Pt(10)
    r1 = p.add_run()
    r1.text = f"{b_label}:\n"
    r1.font.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = NAVY
    r2 = p.add_run()
    r2.text = b_val
    r2.font.size = Pt(12)
    r2.font.color.rgb = DARK_GRAY

# ========================================================
# SLIDE 5: GEOSPATIAL MAP & MULTILINGUAL ADVISORIES
# ========================================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_header(s5, "Interactive Map & Multilingual Advisories", "Sub-District Block Risk Choropleths & Actionable Vernacular Guidance")

map_img = os.path.join(IMG_DIR, "04_geospatial_block_map.png")
if os.path.exists(map_img):
    s5.shapes.add_picture(map_img, Inches(0.8), Inches(1.75), width=Inches(5.7))

adv_img = os.path.join(IMG_DIR, "05_multilingual_advisories_hindi.png")
if os.path.exists(adv_img):
    s5.shapes.add_picture(adv_img, Inches(6.8), Inches(1.75), width=Inches(5.7))

# ========================================================
# SLIDE 6: DISTRICT OFFICER COMMAND CENTER & ANALYTICS
# ========================================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_header(s6, "District Officer & Disaster Management Console", "Multi-Block Monitoring, Threshold Warnings & Automated Alert Dispatching")

off_img = os.path.join(IMG_DIR, "08_officer_command_center.png")
if os.path.exists(off_img):
    s6.shapes.add_picture(off_img, Inches(0.8), Inches(1.75), width=Inches(5.7))

aly_img = os.path.join(IMG_DIR, "09_officer_analytics.png")
if os.path.exists(aly_img):
    s6.shapes.add_picture(aly_img, Inches(6.8), Inches(1.75), width=Inches(5.7))

# ========================================================
# SLIDE 7: BUSINESS MODEL & NATIONAL ROLLOUT PLAN
# ========================================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_header(s7, "National Rollout & Ecosystem Integration", "Scaling from Pilot Districts to Pan-India 6,000+ Administrative Blocks")

road_img = os.path.join(IMG_DIR, "diagram-roadmap.png")
if os.path.exists(road_img):
    s7.shapes.add_picture(road_img, Inches(0.8), Inches(1.75), width=Inches(6.2))

c7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.75), Inches(5.2), Inches(5.1))
c7.fill.solid()
c7.fill.fore_color.rgb = CARD_BG
c7.line.color.rgb = CARD_BORDER
c7.line.width = Pt(1.5)

tf7 = c7.text_frame
tf7.word_wrap = True
tf7.margin_left = tf7.margin_right = Inches(0.3)
tf7.margin_top = Inches(0.3)

pts7 = [
    ("MoES / NCMRWF High-Performance Integration", "Seamless deployment as an operational microservice atop NCMRWF Mihir/Pratyush supercomputing clusters."),
    ("PMFBY Crop Insurance Ground Truth", "Verifiable downscaled hyper-local historical rainfall data for objective crop loss compensation claims."),
    ("KVK & Farmer Producer Organisation (FPO) Reach", "Direct API integration into Meghdoot and Damini apps, with automated WhatsApp/SMS voice alerts in regional dialects."),
    ("Cost Efficiency & Zero Hardware Overhead", "FastAPI + LightGBM architecture runs inference in <120ms with low memory footprint, scaling to all 6,000+ blocks in India at minimal cloud compute cost.")
]

for idx, (t, d) in enumerate(pts7):
    p = tf7.paragraphs[0] if idx == 0 else tf7.add_paragraph()
    p.space_after = Pt(10)
    r1 = p.add_run()
    r1.text = f"• {t}\n"
    r1.font.bold = True
    r1.font.size = Pt(14)
    r1.font.color.rgb = NAVY
    r2 = p.add_run()
    r2.text = d
    r2.font.size = Pt(12)
    r2.font.color.rgb = DARK_GRAY

OUT_PPTX = r"d:\SIH_2026\monsoonguard-system\docs\MonsoonGuard_SIH2026.pptx"
prs.save(OUT_PPTX)
print(f"Presentation saved successfully to {OUT_PPTX}")
